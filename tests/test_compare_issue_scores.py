import csv
import json
from pathlib import Path

from openpyxl import load_workbook
from PIL import Image


ROOT = Path(__file__).resolve().parents[1]


def test_two_judge_comparison_artifacts_are_complete():
    data = json.loads((ROOT / "Issues_No_Tests_scoring_comparison.json").read_text())
    assert data["summary"]["rows"] == 88
    assert len(data["tasks"]) == 88
    assert len({(row["repo"], row["issue_number"]) for row in data["tasks"]}) == 88
    csv_rows = list(csv.DictReader((ROOT / "Issues_No_Tests_scoring_comparison.csv").open()))
    assert len(csv_rows) == 88
    workbook = load_workbook(ROOT / "Issues_No_Tests_scoring_comparison.xlsx", read_only=True)
    assert workbook["Task Comparison"].max_row == 89
    with Image.open(ROOT / "classification_figures" / "7_judge_comparison.png") as image:
        image.verify()


def test_reported_exact_agreement_matches_task_rows():
    data = json.loads((ROOT / "Issues_No_Tests_scoring_comparison.json").read_text())
    rows = data["tasks"]
    summary = data["summary"]
    assert summary["science_share"]["exact_agreement_count"] == sum(
        row["science_pct_delta"] == 0 for row in rows
    )
    assert summary["science_difficulty"]["exact_score_agreement_count"] == sum(
        row["science_difficulty_delta"] == 0 for row in rows
    )
    assert summary["swe_difficulty"]["exact_score_agreement_count"] == sum(
        row["swe_difficulty_delta"] == 0 for row in rows
    )


def test_codex_direct_annotations_are_complete_and_external_api_free():
    rows = json.loads(
        (ROOT / "Issues_No_Tests_final_annotations_codex_subagent.json").read_text()
    )
    assert len(rows) == 88
    assert len({(row["repo"], row["issue_number"]) for row in rows}) == 88
    assert {row["provenance"] for row in rows} == {
        "codex_subagent_direct_no_external_api"
    }
    assert all(row["science_share"] + row["swe_share"] == 100 for row in rows)


def test_three_judge_comparison_artifacts_are_complete():
    report = json.loads(
        (ROOT / "Issues_No_Tests_three_judge_comparison.json").read_text()
    )
    assert report["rows"] == len(report["tasks"]) == 88
    assert set(report["pairwise"]) == {
        "MiniMax run 1 vs MiniMax run 2",
        "MiniMax run 1 vs Codex direct",
        "MiniMax run 2 vs Codex direct",
    }
    workbook = load_workbook(
        ROOT / "Issues_No_Tests_three_judge_comparison.xlsx", read_only=True
    )
    assert workbook["Task Comparison"].max_row == 89
    with Image.open(ROOT / "classification_figures" / "8_three_judge_agreement.png") as image:
        image.verify()
